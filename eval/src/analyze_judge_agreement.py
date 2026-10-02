#!/usr/bin/env python3
"""Analyse agreement between two or more model judges.

The script deliberately separates absolute agreement from correlation:

* ICC(A,1) answers whether judges give nearly the same numerical scores.
* ICC(C,1), Pearson, Spearman, and Kendall's W answer whether judges move or
  rank samples similarly after systematic severity differences are ignored.

Example:

    python eval/src/analyze_judge_agreement.py \
      --judge Judge_A=eval/eval_using_understanding/scores_summary.xlsx \
      --judge Gemini_2.5_Pro=eval/eval_using_understanding-gemini25p/scores_summary.xlsx \
      --output-dir eval/judge_agreement/understanding_vs_gemini25p
"""

from __future__ import annotations

import argparse
import csv
import json
import math
from dataclasses import dataclass
from itertools import combinations
from pathlib import Path
from typing import Any, Callable, Iterable, Sequence

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


DEFAULT_SCORE_COLUMNS = [
    "modeling_rigor",
    "practicality_scientificity",
    "problem_analysis",
    "result_bias",
]


@dataclass(frozen=True)
class JudgeInput:
    label: str
    path: Path


def parse_judge_spec(value: str) -> JudgeInput:
    """Parse ``LABEL=PATH`` while allowing ``=`` inside the path."""
    if "=" not in value:
        raise argparse.ArgumentTypeError("--judge 必须使用 LABEL=PATH 格式")
    label, raw_path = value.split("=", 1)
    label = label.strip()
    raw_path = raw_path.strip()
    if not label or not raw_path:
        raise argparse.ArgumentTypeError("--judge 的 LABEL 和 PATH 均不能为空")
    return JudgeInput(label=label, path=Path(raw_path))


def _normalise_key(value: Any) -> str:
    if value is None:
        raise ValueError("key value is empty")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _as_finite_float(value: Any, *, source: str) -> float:
    if isinstance(value, bool):
        raise ValueError(f"{source} is boolean, not a score")
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{source} is not numeric: {value!r}") from exc
    if not math.isfinite(result):
        raise ValueError(f"{source} is not finite: {value!r}")
    return result


def load_score_workbook(
    path: Path,
    sheet_name: str,
    key_columns: Sequence[str],
    score_columns: Sequence[str],
) -> dict[tuple[str, ...], dict[str, float]]:
    """Load and validate one judge's task-level score workbook."""
    if not path.is_file():
        raise FileNotFoundError(f"评分文件不存在: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(
            f"{path} 中没有 sheet {sheet_name!r}; 可用 sheets: {workbook.sheetnames}"
        )

    worksheet = workbook[sheet_name]
    rows = worksheet.iter_rows(values_only=True)
    raw_header = next(rows, None)
    if raw_header is None:
        raise ValueError(f"{path}:{sheet_name} 是空表")
    header = [str(value).strip() if value is not None else "" for value in raw_header]
    if len(set(header)) != len(header):
        raise ValueError(f"{path}:{sheet_name} 存在重复列名")

    required = list(key_columns) + list(score_columns)
    missing = [column for column in required if column not in header]
    if missing:
        raise ValueError(f"{path}:{sheet_name} 缺少列: {missing}")

    result: dict[tuple[str, ...], dict[str, float]] = {}
    for row_number, values in enumerate(rows, start=2):
        if not values or all(value is None for value in values):
            continue
        record = dict(zip(header, values))
        try:
            key = tuple(_normalise_key(record[column]) for column in key_columns)
        except ValueError as exc:
            raise ValueError(f"{path}:{sheet_name} 第 {row_number} 行 key 无效: {exc}") from exc
        if key in result:
            raise ValueError(f"{path}:{sheet_name} 存在重复 key: {key}")
        result[key] = {
            column: _as_finite_float(
                record[column], source=f"{path}:{sheet_name} 第 {row_number} 行 {column}"
            )
            for column in score_columns
        }

    if not result:
        raise ValueError(f"{path}:{sheet_name} 没有有效数据行")
    return result


def rankdata(values: Sequence[float]) -> np.ndarray:
    """Return average ranks (1-based), including correct handling of ties."""
    array = np.asarray(values, dtype=float)
    order = np.argsort(array, kind="mergesort")
    ranks = np.empty(len(array), dtype=float)
    start = 0
    while start < len(array):
        end = start + 1
        while end < len(array) and array[order[end]] == array[order[start]]:
            end += 1
        ranks[order[start:end]] = (start + end - 1) / 2.0 + 1.0
        start = end
    return ranks


def pearson_correlation(x: Sequence[float], y: Sequence[float]) -> float:
    x_array = np.asarray(x, dtype=float)
    y_array = np.asarray(y, dtype=float)
    x_centered = x_array - x_array.mean()
    y_centered = y_array - y_array.mean()
    denominator = math.sqrt(float(x_centered @ x_centered) * float(y_centered @ y_centered))
    if denominator == 0:
        return math.nan
    return float((x_centered @ y_centered) / denominator)


def spearman_correlation(x: Sequence[float], y: Sequence[float]) -> float:
    return pearson_correlation(rankdata(x), rankdata(y))


def lin_concordance_correlation(x: Sequence[float], y: Sequence[float]) -> float:
    """Lin's CCC, which penalises location and scale differences."""
    x_array = np.asarray(x, dtype=float)
    y_array = np.asarray(y, dtype=float)
    x_variance = float(np.var(x_array))
    y_variance = float(np.var(y_array))
    covariance = float(np.mean((x_array - x_array.mean()) * (y_array - y_array.mean())))
    denominator = x_variance + y_variance + float((x_array.mean() - y_array.mean()) ** 2)
    if denominator == 0:
        return math.nan
    return 2.0 * covariance / denominator


def intraclass_correlations(score_matrix: np.ndarray) -> tuple[float, float]:
    """Return ICC(A,1) and ICC(C,1) from a complete target-by-judge matrix.

    The formulas follow the two-way random-effects, single-measure absolute
    agreement ICC(A,1), and the two-way mixed-effects, single-measure
    consistency ICC(C,1). The column mean-square term is required for the
    absolute-agreement statistic.
    """
    matrix = np.asarray(score_matrix, dtype=float)
    if matrix.ndim != 2:
        raise ValueError("score_matrix must be two-dimensional")
    target_count, judge_count = matrix.shape
    if target_count < 2 or judge_count < 2:
        return math.nan, math.nan

    target_means = matrix.mean(axis=1)
    judge_means = matrix.mean(axis=0)
    grand_mean = float(matrix.mean())

    ss_targets = judge_count * float(np.sum((target_means - grand_mean) ** 2))
    ss_judges = target_count * float(np.sum((judge_means - grand_mean) ** 2))
    residuals = matrix - target_means[:, None] - judge_means[None, :] + grand_mean
    ss_error = float(np.sum(residuals**2))

    ms_targets = ss_targets / (target_count - 1)
    ms_judges = ss_judges / (judge_count - 1)
    ms_error = ss_error / ((target_count - 1) * (judge_count - 1))

    absolute_denominator = (
        ms_targets
        + (judge_count - 1) * ms_error
        + judge_count * (ms_judges - ms_error) / target_count
    )
    consistency_denominator = ms_targets + (judge_count - 1) * ms_error

    icc_absolute = (
        (ms_targets - ms_error) / absolute_denominator
        if absolute_denominator != 0
        else math.nan
    )
    icc_consistency = (
        (ms_targets - ms_error) / consistency_denominator
        if consistency_denominator != 0
        else math.nan
    )
    return float(icc_absolute), float(icc_consistency)


def kendalls_w(score_matrix: np.ndarray) -> float:
    """Kendall's coefficient of concordance with tie correction."""
    matrix = np.asarray(score_matrix, dtype=float)
    if matrix.ndim != 2:
        raise ValueError("score_matrix must be two-dimensional")
    target_count, judge_count = matrix.shape
    if target_count < 2 or judge_count < 2:
        return math.nan

    ranked = np.column_stack([rankdata(matrix[:, index]) for index in range(judge_count)])
    rank_sums = ranked.sum(axis=1)
    expected_rank_sum = judge_count * (target_count + 1) / 2.0
    squared_deviation = float(np.sum((rank_sums - expected_rank_sum) ** 2))

    tie_correction = 0.0
    for judge_index in range(judge_count):
        _, tie_counts = np.unique(matrix[:, judge_index], return_counts=True)
        tie_correction += float(np.sum(tie_counts**3 - tie_counts))

    denominator = (
        judge_count**2 * (target_count**3 - target_count)
        - judge_count * tie_correction
    )
    if denominator == 0:
        return math.nan
    return 12.0 * squared_deviation / denominator


def _finite_float(value: float) -> float | None:
    return float(value) if math.isfinite(float(value)) else None


def _metric_interval(values: Iterable[float], confidence: float) -> dict[str, float | None]:
    finite = np.asarray([value for value in values if math.isfinite(float(value))], dtype=float)
    if finite.size == 0:
        return {"lower": None, "upper": None}
    alpha = 1.0 - confidence
    lower, upper = np.quantile(finite, [alpha / 2.0, 1.0 - alpha / 2.0])
    return {"lower": float(lower), "upper": float(upper)}


def agreement_level(value: float | None) -> str:
    """Conventional ICC interpretation from Koo and Li (2016)."""
    if value is None or not math.isfinite(value):
        return "undefined"
    if value >= 0.90:
        return "excellent"
    if value >= 0.75:
        return "good"
    if value >= 0.50:
        return "moderate"
    return "poor"


def _pair_metrics(x: np.ndarray, y: np.ndarray) -> dict[str, float]:
    difference = y - x
    return {
        "pearson_r": pearson_correlation(x, y),
        "spearman_rho": spearman_correlation(x, y),
        "lin_ccc": lin_concordance_correlation(x, y),
        "mean_bias_b_minus_a": float(difference.mean()),
        "mae": float(np.mean(np.abs(difference))),
        "rmse": float(math.sqrt(float(np.mean(difference**2)))),
        "within_0_5": float(np.mean(np.abs(difference) <= 0.5)),
        "within_1_0": float(np.mean(np.abs(difference) <= 1.0)),
        "bland_altman_lower": float(difference.mean() - 1.96 * difference.std(ddof=1)),
        "bland_altman_upper": float(difference.mean() + 1.96 * difference.std(ddof=1)),
    }


def analyse_score_matrix(
    score_matrix: np.ndarray,
    judge_labels: Sequence[str],
    *,
    bootstrap_samples: int,
    confidence: float,
    rng: np.random.Generator,
    excellent_threshold: float,
    bootstrap_groups: Sequence[str] | None = None,
    score_maximum: float = 10.0,
) -> dict[str, Any]:
    target_count, judge_count = score_matrix.shape
    if bootstrap_groups is not None and len(bootstrap_groups) != target_count:
        raise ValueError("bootstrap_groups length must match score_matrix rows")
    icc_absolute, icc_consistency = intraclass_correlations(score_matrix)
    concordance = kendalls_w(score_matrix)

    descriptions = {}
    for index, label in enumerate(judge_labels):
        values = score_matrix[:, index]
        descriptions[label] = {
            "mean": float(values.mean()),
            "sd": float(values.std(ddof=1)),
            "min": float(values.min()),
            "max": float(values.max()),
            "at_maximum_rate": float(np.mean(values == score_maximum)),
        }

    pair_results: list[dict[str, Any]] = []
    for first, second in combinations(range(judge_count), 2):
        metrics = _pair_metrics(score_matrix[:, first], score_matrix[:, second])
        pair_results.append(
            {
                "judge_a": judge_labels[first],
                "judge_b": judge_labels[second],
                "mean_a": descriptions[judge_labels[first]]["mean"],
                "mean_b": descriptions[judge_labels[second]]["mean"],
                "at_maximum_rate_a": descriptions[judge_labels[first]]["at_maximum_rate"],
                "at_maximum_rate_b": descriptions[judge_labels[second]]["at_maximum_rate"],
                **{key: _finite_float(value) for key, value in metrics.items()},
                "_indices": (first, second),
            }
        )

    global_bootstrap: dict[str, list[float]] = {
        "icc_absolute_a1": [],
        "icc_consistency_c1": [],
        "kendall_w": [],
    }
    pair_bootstrap: list[dict[str, list[float]]] = [
        {"pearson_r": [], "spearman_rho": [], "lin_ccc": []} for _ in pair_results
    ]

    group_indices: list[np.ndarray] | None = None
    if bootstrap_groups is not None:
        unique_groups = list(dict.fromkeys(bootstrap_groups))
        group_indices = [
            np.asarray(
                [index for index, group in enumerate(bootstrap_groups) if group == value],
                dtype=int,
            )
            for value in unique_groups
        ]

    for _ in range(bootstrap_samples):
        if group_indices is None:
            sample_indices = rng.integers(0, target_count, size=target_count)
        else:
            selected_groups = rng.integers(0, len(group_indices), size=len(group_indices))
            sample_indices = np.concatenate([group_indices[index] for index in selected_groups])
        sample = score_matrix[sample_indices]
        absolute_sample, consistency_sample = intraclass_correlations(sample)
        global_bootstrap["icc_absolute_a1"].append(absolute_sample)
        global_bootstrap["icc_consistency_c1"].append(consistency_sample)
        global_bootstrap["kendall_w"].append(kendalls_w(sample))
        for pair_index, pair_result in enumerate(pair_results):
            first, second = pair_result["_indices"]
            boot_metrics = _pair_metrics(sample[:, first], sample[:, second])
            for metric_name in pair_bootstrap[pair_index]:
                pair_bootstrap[pair_index][metric_name].append(boot_metrics[metric_name])

    global_metrics = {
        "icc_absolute_a1": {
            "value": _finite_float(icc_absolute),
            "ci": _metric_interval(global_bootstrap["icc_absolute_a1"], confidence),
            "level": agreement_level(_finite_float(icc_absolute)),
        },
        "icc_consistency_c1": {
            "value": _finite_float(icc_consistency),
            "ci": _metric_interval(global_bootstrap["icc_consistency_c1"], confidence),
            "level": agreement_level(_finite_float(icc_consistency)),
        },
        "kendall_w": {
            "value": _finite_float(concordance),
            "ci": _metric_interval(global_bootstrap["kendall_w"], confidence),
        },
    }

    for pair_index, pair_result in enumerate(pair_results):
        pair_result["confidence_intervals"] = {
            metric_name: _metric_interval(values, confidence)
            for metric_name, values in pair_bootstrap[pair_index].items()
        }
        del pair_result["_indices"]

    lower_bound = global_metrics["icc_absolute_a1"]["ci"]["lower"]
    return {
        "target_count": target_count,
        "descriptive": descriptions,
        "global": global_metrics,
        "pairwise": pair_results,
        "decision": {
            "threshold": excellent_threshold,
            "point_estimate_pass": bool(
                math.isfinite(icc_absolute) and icc_absolute >= excellent_threshold
            ),
            "confidence_bound_pass": bool(
                lower_bound is not None and lower_bound >= excellent_threshold
            ),
        },
    }


def _fmt(value: float | None, digits: int = 3) -> str:
    return "NA" if value is None else f"{value:.{digits}f}"


def _fmt_ci(metric: dict[str, Any]) -> str:
    value = metric["value"]
    interval = metric["ci"]
    return f"{_fmt(value)} [{_fmt(interval['lower'])}, {_fmt(interval['upper'])}]"


def build_markdown_report(results: dict[str, Any]) -> str:
    overall = results["outcomes"]["overall_mean"]
    absolute = overall["global"]["icc_absolute_a1"]
    consistency = overall["global"]["icc_consistency_c1"]
    threshold = results["configuration"]["excellent_threshold"]
    confidence_label = f"{results['configuration']['confidence']:.0%} CI"

    if overall["decision"]["confidence_bound_pass"]:
        conclusion = (
            f"总体平均分的 ICC(A,1) 置信区间下界达到 {threshold:.2f}，"
            "可视为对“高度绝对一致”的强证据。"
        )
    elif overall["decision"]["point_estimate_pass"]:
        conclusion = (
            f"总体平均分的 ICC(A,1) 点估计达到 {threshold:.2f}，但置信区间下界未达到；"
            "证据尚不足以支持稳定的高度绝对一致。"
        )
    else:
        conclusion = (
            f"总体平均分的 ICC(A,1) 未达到预设阈值 {threshold:.2f}，"
            "当前数据不支持“不同裁判给出高度一致的绝对分数”。"
        )

    lines = [
        "# 模型裁判一致性实验报告",
        "",
        "## 结论",
        "",
        conclusion,
        "",
        f"- 完全配对样本数：{results['coverage']['common_target_count']}",
        f"- 总体 ICC(A,1)：{_fmt_ci(absolute)}（{absolute['level']}）",
        f"- 总体 ICC(C,1)：{_fmt_ci(consistency)}（{consistency['level']}）",
        "- ICC(A,1) 衡量绝对分数是否一致；ICC(C,1) 仅衡量去除系统性宽严差后的变化一致性。",
        "",
        "## 数据覆盖",
        "",
        "| 裁判 | 原始样本数 | 未进入共同交集 | 来源 |",
        "|---|---:|---:|---|",
    ]
    for judge in results["judges"]:
        lines.append(
            f"| {judge['label']} | {judge['row_count']} | {judge['excluded_from_intersection']} | `{judge['path']}` |"
        )

    lines.extend(
        [
            "",
            "## 各维度一致性",
            "",
            f"| 评分项 | ICC(A,1), {confidence_label} | ICC(C,1), {confidence_label} | Kendall W, {confidence_label} | 绝对一致等级 |",
            "|---|---:|---:|---:|---|",
        ]
    )
    outcome_order = ["overall_mean", *results["configuration"]["score_columns"]]
    for outcome_name in outcome_order:
        outcome = results["outcomes"][outcome_name]
        global_metrics = outcome["global"]
        lines.append(
            "| "
            + " | ".join(
                [
                    outcome_name,
                    _fmt_ci(global_metrics["icc_absolute_a1"]),
                    _fmt_ci(global_metrics["icc_consistency_c1"]),
                    _fmt_ci(global_metrics["kendall_w"]),
                    global_metrics["icc_absolute_a1"]["level"],
                ]
            )
            + " |"
        )

    lines.extend(
        [
            "",
            "## 总体平均分的两两诊断",
            "",
            "| 裁判 A | 裁判 B | 均值 A | 均值 B | 满分率 A | 满分率 B | B-A 偏差 | MAE | Pearson | Spearman | Lin CCC | |差|≤1 |",
            "|---|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|",
        ]
    )
    for pair in overall["pairwise"]:
        lines.append(
            "| "
            + " | ".join(
                [
                    pair["judge_a"],
                    pair["judge_b"],
                    _fmt(pair["mean_a"]),
                    _fmt(pair["mean_b"]),
                    f"{pair['at_maximum_rate_a']:.1%}",
                    f"{pair['at_maximum_rate_b']:.1%}",
                    _fmt(pair["mean_bias_b_minus_a"]),
                    _fmt(pair["mae"]),
                    _fmt(pair["pearson_r"]),
                    _fmt(pair["spearman_rho"]),
                    _fmt(pair["lin_ccc"]),
                    _fmt(pair["within_1_0"]),
                ]
            )
            + " |"
        )

    lines.extend(
        [
            "",
            "## 判定规则",
            "",
            f"- 预注册主判据：总体平均分 ICC(A,1) ≥ {threshold:.2f}。",
            f"- 强证据判据：ICC(A,1) 的 {results['configuration']['confidence']:.0%} bootstrap 置信区间下界也 ≥ {threshold:.2f}。",
            "- 相关系数高但 ICC(A,1) 低，表示排序可能相近，但原始分数不能互换。",
            "- 两位裁判时 Kendall W 与 Spearman 相关存在直接对应关系，不能把 W 的数值直接当成绝对一致性。",
            (
                f"- Bootstrap 以 `{results['configuration']['bootstrap_cluster_column']}` 为聚类单位重采样，"
                "保持同一题目的多个模型回答和所有裁判评分一起出现。"
                if results["configuration"]["bootstrap_cluster_column"]
                else "- Bootstrap 以被评分样本为单位重采样，保持每个样本的所有裁判评分成对出现。"
            ),
            "",
        ]
    )
    return "\n".join(lines)


def _write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: Sequence[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def _style_worksheet(worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_cells in worksheet.columns:
        width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(width + 2, 55)


def write_excel_summary(
    path: Path,
    results: dict[str, Any],
    multi_rows: list[dict[str, Any]],
    pair_rows: list[dict[str, Any]],
    paired_rows: list[dict[str, Any]],
    multi_fields: Sequence[str],
    pair_fields: Sequence[str],
    paired_fields: Sequence[str],
) -> None:
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    sheets = [
        ("Multi_Rater", multi_fields, multi_rows),
        ("Pairwise", pair_fields, pair_rows),
        ("Paired_Scores", paired_fields, paired_rows),
    ]
    for sheet_name, fields, rows in sheets:
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(list(fields))
        for row in rows:
            worksheet.append(
                [
                    json.dumps(row.get(field), ensure_ascii=False, sort_keys=True)
                    if isinstance(row.get(field), (dict, list))
                    else row.get(field)
                    for field in fields
                ]
            )
        _style_worksheet(worksheet)

    metadata = workbook.create_sheet("Metadata")
    metadata.append(["field", "value"])
    metadata.append(["common_target_count", results["coverage"]["common_target_count"]])
    for key, value in results["configuration"].items():
        metadata.append(
            [key, json.dumps(value, ensure_ascii=False) if isinstance(value, list) else value]
        )
    for judge in results["judges"]:
        metadata.append([f"judge:{judge['label']}", judge["path"]])
    _style_worksheet(metadata)
    workbook.save(path)


def write_outputs(
    results: dict[str, Any],
    output_dir: Path,
    common_keys: Sequence[tuple[str, ...]],
    aligned: dict[str, dict[tuple[str, ...], dict[str, float]]],
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)

    with (output_dir / "agreement_results.json").open("w", encoding="utf-8") as handle:
        json.dump(results, handle, ensure_ascii=False, indent=2, allow_nan=False)

    (output_dir / "report.md").write_text(build_markdown_report(results), encoding="utf-8")

    multi_rows: list[dict[str, Any]] = []
    pair_rows: list[dict[str, Any]] = []
    for outcome_name, outcome in results["outcomes"].items():
        global_metrics = outcome["global"]
        row: dict[str, Any] = {"outcome": outcome_name, "n": outcome["target_count"]}
        for metric_name in ["icc_absolute_a1", "icc_consistency_c1", "kendall_w"]:
            metric = global_metrics[metric_name]
            row[metric_name] = metric["value"]
            row[f"{metric_name}_ci_lower"] = metric["ci"]["lower"]
            row[f"{metric_name}_ci_upper"] = metric["ci"]["upper"]
        row["absolute_agreement_level"] = global_metrics["icc_absolute_a1"]["level"]
        row["point_estimate_pass"] = outcome["decision"]["point_estimate_pass"]
        row["confidence_bound_pass"] = outcome["decision"]["confidence_bound_pass"]
        multi_rows.append(row)

        for pair in outcome["pairwise"]:
            pair_rows.append({"outcome": outcome_name, **pair})

    multi_fields = [
        "outcome",
        "n",
        "icc_absolute_a1",
        "icc_absolute_a1_ci_lower",
        "icc_absolute_a1_ci_upper",
        "icc_consistency_c1",
        "icc_consistency_c1_ci_lower",
        "icc_consistency_c1_ci_upper",
        "kendall_w",
        "kendall_w_ci_lower",
        "kendall_w_ci_upper",
        "absolute_agreement_level",
        "point_estimate_pass",
        "confidence_bound_pass",
    ]
    _write_csv(output_dir / "multi_rater_metrics.csv", multi_rows, multi_fields)

    pair_fields = [
        "outcome",
        "judge_a",
        "judge_b",
        "mean_a",
        "mean_b",
        "at_maximum_rate_a",
        "at_maximum_rate_b",
        "pearson_r",
        "spearman_rho",
        "lin_ccc",
        "mean_bias_b_minus_a",
        "mae",
        "rmse",
        "within_0_5",
        "within_1_0",
        "bland_altman_lower",
        "bland_altman_upper",
        "confidence_intervals",
    ]
    pair_csv_rows = []
    for row in pair_rows:
        csv_row = dict(row)
        csv_row["confidence_intervals"] = json.dumps(
            csv_row["confidence_intervals"], ensure_ascii=False, sort_keys=True
        )
        pair_csv_rows.append(csv_row)
    _write_csv(output_dir / "pairwise_metrics.csv", pair_csv_rows, pair_fields)

    key_columns = results["configuration"]["key_columns"]
    score_columns = results["configuration"]["score_columns"]
    judge_labels = [judge["label"] for judge in results["judges"]]
    paired_fields = list(key_columns)
    for label in judge_labels:
        paired_fields.extend([f"{label}:{column}" for column in score_columns])
        paired_fields.append(f"{label}:overall_mean")

    paired_rows = []
    for key in common_keys:
        row = dict(zip(key_columns, key))
        for label in judge_labels:
            values = aligned[label][key]
            for column in score_columns:
                row[f"{label}:{column}"] = values[column]
            row[f"{label}:overall_mean"] = float(
                np.mean([values[column] for column in score_columns])
            )
        paired_rows.append(row)
    _write_csv(output_dir / "paired_scores.csv", paired_rows, paired_fields)
    write_excel_summary(
        output_dir / "agreement_summary.xlsx",
        results,
        multi_rows,
        pair_rows,
        paired_rows,
        multi_fields,
        pair_fields,
        paired_fields,
    )


def run_analysis(args: argparse.Namespace) -> dict[str, Any]:
    if len(args.judge) < 2:
        raise ValueError("至少需要两个 --judge LABEL=PATH")
    labels = [judge.label for judge in args.judge]
    if len(labels) != len(set(labels)):
        raise ValueError("--judge LABEL 必须唯一")
    if not 0.0 < args.confidence < 1.0:
        raise ValueError("--confidence 必须在 0 和 1 之间")
    if args.bootstrap < 0:
        raise ValueError("--bootstrap 不能为负数")

    loaded = {
        judge.label: load_score_workbook(
            judge.path,
            args.sheet,
            args.key_columns,
            args.score_columns,
        )
        for judge in args.judge
    }
    common_key_set = set.intersection(*(set(rows) for rows in loaded.values()))
    common_keys = sorted(common_key_set)
    if len(common_keys) < 2:
        raise ValueError(f"共同有效样本不足 2 个，实际为 {len(common_keys)} 个")

    rng = np.random.default_rng(args.seed)
    bootstrap_groups = None
    if args.bootstrap_cluster_column:
        if args.bootstrap_cluster_column not in args.key_columns:
            raise ValueError(
                "--bootstrap-cluster-column 必须同时出现在 --key-columns 中"
            )
        cluster_index = args.key_columns.index(args.bootstrap_cluster_column)
        bootstrap_groups = [key[cluster_index] for key in common_keys]
    outcomes: dict[str, Any] = {}
    outcome_columns: list[tuple[str, Callable[[dict[str, float]], float]]] = [
        (
            "overall_mean",
            lambda row: float(np.mean([row[column] for column in args.score_columns])),
        )
    ]
    outcome_columns.extend((column, lambda row, c=column: row[c]) for column in args.score_columns)

    for outcome_name, extractor in outcome_columns:
        matrix = np.asarray(
            [
                [extractor(loaded[label][key]) for label in labels]
                for key in common_keys
            ],
            dtype=float,
        )
        outcomes[outcome_name] = analyse_score_matrix(
            matrix,
            labels,
            bootstrap_samples=args.bootstrap,
            confidence=args.confidence,
            rng=rng,
            excellent_threshold=args.excellent_threshold,
            bootstrap_groups=bootstrap_groups,
            score_maximum=args.score_maximum,
        )

    results = {
        "configuration": {
            "sheet": args.sheet,
            "key_columns": list(args.key_columns),
            "score_columns": list(args.score_columns),
            "bootstrap_samples": args.bootstrap,
            "bootstrap_cluster_column": args.bootstrap_cluster_column,
            "confidence": args.confidence,
            "seed": args.seed,
            "excellent_threshold": args.excellent_threshold,
            "score_maximum": args.score_maximum,
        },
        "coverage": {
            "common_target_count": len(common_keys),
            "intersection_policy": "complete-case inner join on all key columns",
            "bootstrap_cluster_count": (
                len(set(bootstrap_groups)) if bootstrap_groups is not None else None
            ),
        },
        "judges": [
            {
                "label": judge.label,
                "path": str(judge.path),
                "row_count": len(loaded[judge.label]),
                "excluded_from_intersection": len(loaded[judge.label]) - len(common_keys),
            }
            for judge in args.judge
        ],
        "outcomes": outcomes,
    }
    write_outputs(results, args.output_dir, common_keys, loaded)
    return results


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="比较两个或多个模型裁判在同一批样本上的评分一致性"
    )
    parser.add_argument(
        "--judge",
        action="append",
        type=parse_judge_spec,
        required=True,
        help="裁判评分文件，格式 LABEL=PATH；至少传两次",
    )
    parser.add_argument("--sheet", default="Task_Stage_Scores")
    parser.add_argument(
        "--key-columns", nargs="+", default=["Model Name", "Task ID"]
    )
    parser.add_argument(
        "--score-columns", nargs="+", default=DEFAULT_SCORE_COLUMNS
    )
    parser.add_argument("--bootstrap", type=int, default=2000)
    parser.add_argument(
        "--bootstrap-cluster-column",
        default=None,
        help="可选：按 key 中的一列做聚类 bootstrap，例如 Task ID",
    )
    parser.add_argument("--confidence", type=float, default=0.95)
    parser.add_argument("--seed", type=int, default=20261001)
    parser.add_argument("--excellent-threshold", type=float, default=0.90)
    parser.add_argument("--score-maximum", type=float, default=10.0)
    parser.add_argument("--output-dir", type=Path, required=True)
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    results = run_analysis(args)
    overall = results["outcomes"]["overall_mean"]
    absolute = overall["global"]["icc_absolute_a1"]
    consistency = overall["global"]["icc_consistency_c1"]
    print(f"Matched targets: {results['coverage']['common_target_count']}")
    print(f"ICC(A,1): {_fmt_ci(absolute)} ({absolute['level']})")
    print(f"ICC(C,1): {_fmt_ci(consistency)} ({consistency['level']})")
    print(f"High absolute agreement pass: {overall['decision']['point_estimate_pass']}")
    print(f"Report: {args.output_dir / 'report.md'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
